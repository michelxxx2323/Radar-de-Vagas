#!/usr/bin/env python3
"""
Kit de candidatura: CV adaptado em PDF + carta + respostas prontas para uma vaga.

Pega uma vaga do último relatório (output/vagas_triadas_*.csv), junta com a
descrição completa (output/vagas_coletadas.json) e pede ao Claude um CV
montado SÓ com o que está no cv.txt e no regras-cv.md (que também traz as
regras de formato, o que ele não tem e o banco de resultados). Em duas etapas:

1. checagem: requisitos OBRIGATÓRIOS da vaga que o currículo não atende. Se
   houver, pergunta "montar mesmo assim?" antes de gastar com o kit;
2. kit: CV, carta, respostas para o formulário.

O CV vira .docx no template dos CVs anteriores e o Word converte pra PDF. O
modelo escreve bullets de sobra, do mais pro menos relevante, e o script corta
do fim até o PDF ter exatamente 2 páginas. Checagens determinísticas apontam
números que não estão no cv.txt/regras e termos que as regras proíbem.

Sai em output/kits/<empresa>_<vaga>/ (docx, PDF, kit.md, kit.json); o PDF
também é copiado pra pasta `kit.pasta_pdf` do config.yaml (OneDrive), e cada
vaga entra em output/kits/registro.csv.

Uso:
    python kit.py                  # lista as vagas "aplicar_agora" do último relatório
    python kit.py --busca amazon   # filtra a lista por empresa/título
    python kit.py --todas          # lista todas as recomendações, não só "aplicar_agora"
    python kit.py 3 8              # gera o kit das vagas 3 e 8 da lista
    python kit.py --link URL       # gera o kit da vaga com esse link
    python kit.py 3 --refazer      # gera de novo, mesmo se o kit já existir
    python kit.py 3 --sempre       # monta mesmo se faltar requisito obrigatório, sem perguntar
    python kit.py 3 --mock         # sem API: testa o fluxo com um kit simulado

Precisa de ANTHROPIC_API_KEY (no .env ou no ambiente) — ver .env.example —
e do Word instalado pra gerar o PDF.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
import shutil
import sys
import unicodedata
from pathlib import Path

import requests
import yaml

from candidaturas import POS_ENVIO, Candidaturas, alertas_do_kit
from novidades import _ler_env
from screen import carregar_preferencias

# preço por milhão de tokens (out/2026), só pra mostrar quanto cada kit custou
PRECO_MTOK = {"claude-opus-5-5": {"entrada": 4.00, "saida": 20.00, "cache_leitura": 0.20, "cache_escrita": 5.00}}
MAX_NA_LISTA = 30
# a busca do beBee só traz ~300 caracteres de cada vaga; abaixo disso, o kit
# busca a página da vaga antes de gerar
DESCRICAO_CURTA = 1000
# Times New Roman 11.5pt negrito na largura útil do A4 com as margens do
# template: acima disso o título quebra em duas linhas
TITULO_MAX_CHARS = 78
PAGINAS = 2
REGRAS_PATH = Path("regras-cv.md")

INSTRUCOES = """\
Você monta CVs e kits de candidatura para o candidato abaixo. Há duas fontes de
verdade: o currículo-mestre (<curriculo_mestre>) e as regras dele (<regras>), que
também trazem fatos confirmados (ferramentas, banco de resultados) e o que ele NÃO
tem. Nada fora dessas duas fontes entra no CV, na carta ou nas respostas: se a vaga
pede algo que ele não comprova, fica de fora. Onde as regras e estas instruções
divergirem, valem as regras.

Você roda dentro de um script, sem conversa com o candidato. O script cuida de
arquivos, PDF, contagem de páginas, pasta de destino, nome do arquivo e registro da
vaga — ignore essas partes das regras. O que dependeria de confirmar com ele fica de
fora.

Para o CV caber em exatamente 2 páginas, o script corta bullets do fim de cada
experiência. Por isso escreva bullets de sobra — 6 a 8 por experiência, sem contar a
linha de resultados — sempre do mais para o menos relevante para a vaga.

Campos:
- idioma: o da descrição da vaga ("pt" ou "en"); tudo é escrito nele, números iguais.
- variante: "financeira" só quando a empresa é instituição financeira (regras, seção
  4); senão "padrao".
- titulo_cv: uma linha de até 75 caracteres, frentes separadas por " | ", alinhadas
  à vaga e só com funções que ele exerceu.
- resumo: um parágrafo de 120 a 160 palavras, na primeira pessoa ("desenvolvi",
  "construí"), com a abertura que as regras indicarem.
- experiencias: todas as do mestre, na ordem das regras. empresa e cargo como no
  mestre (a área depois do "|" pode ser encurtada, nunca trocada por outra função);
  periodo no formato "10/2024 – 01/2026"; local como no mestre; linha_extra é uma
  linha de destaque logo abaixo do período, só onde as regras pedirem (vazia nas
  outras); bullets na primeira pessoa do passado ("Estruturei...", em inglês
  "Built..."); resultados é o texto da linha "Resultados:", sem o prefixo, tirado do
  banco de resultados.
- formacao: um item por curso (titulo em negrito; linhas em itálico embaixo),
  seguindo as regras de formação.
- certificacoes: só na variante financeira; vazio na padrão.
- competencias_chave: 6 itens curtos, alinhados à vaga, todos comprovados.
- ferramentas: 3 linhas agrupadas por tipo, itens separados por " · ".
- idiomas: um item por idioma, no formato "Idioma: Nível", como no mestre.
- carta: 180 a 250 palavras, primeira pessoa, direta, sem clichês ("venho por meio
  desta", "sou apaixonado por"), ligando 2 ou 3 resultados concretos ao que a vaga pede.
- respostas: 5 a 7 perguntas que formulários de triagem costumam fazer para esta vaga
  (por que a empresa, experiência com o que a vaga pede, um resultado de que se
  orgulha, pretensão salarial, disponibilidade...). Quando a resposta depende de algo
  que não está nas fontes (pretensão, data de início), deixe um marcador entre
  colchetes, como [sua pretensão salarial], em vez de inventar.
- palavras_chave: termos da vaga que o CV cobre.
- As preferências dizem o que ele quer e não quer fazer: use-as para escolher o
  ângulo, nunca como conteúdo.
- Nome e contatos já estão no cabeçalho do documento; não os repita.
"""

PEDIDO_CHECAGEM = """\
Antes do kit: liste só os requisitos OBRIGATÓRIOS desta vaga que ele não atende,
seguindo "O que conta como lacuna e o que não conta" nas regras. Diferenciais e
desejáveis nunca entram; viagem e mudança nunca entram; se uma alternativa com "ou"
é atendida, o requisito está atendido. Se ele atende todos, devolva a lista vazia."""

_ITEM = {"type": "string"}
_LISTA = {"type": "array", "items": _ITEM}


def _objeto(props: dict) -> dict:
    return {"type": "object", "properties": props, "required": list(props), "additionalProperties": False}


CHECAGEM_SCHEMA = _objeto({
    "obrigatorios_faltando": {"type": "array", "items": _objeto({"requisito": _ITEM, "motivo": _ITEM})},
})

KIT_SCHEMA = _objeto({
    "idioma": {"type": "string", "enum": ["pt", "en"]},
    "variante": {"type": "string", "enum": ["padrao", "financeira"]},
    "titulo_cv": _ITEM,
    "resumo": _ITEM,
    "experiencias": {"type": "array", "items": _objeto({
        "empresa": _ITEM, "cargo": _ITEM, "periodo": _ITEM, "local": _ITEM, "linha_extra": _ITEM,
        "bullets": _LISTA, "resultados": _ITEM,
    })},
    "formacao": {"type": "array", "items": _objeto({"titulo": _ITEM, "linhas": _LISTA})},
    "certificacoes": {"type": "array", "items": _objeto({"titulo": _ITEM, "linhas": _LISTA})},
    "competencias_chave": _LISTA,
    "ferramentas": _LISTA,
    "idiomas": _LISTA,
    "carta": _ITEM,
    "respostas": {"type": "array", "items": _objeto({"pergunta": _ITEM, "resposta": _ITEM})},
    "palavras_chave": _LISTA,
})

ROTULOS = {
    "pt": {"resumo": "RESUMO PROFISSIONAL", "exp": "EXPERIÊNCIA PROFISSIONAL", "form": "FORMAÇÃO ACADÊMICA",
           "cert": "CERTIFICAÇÕES", "comp": "COMPETÊNCIAS & FERRAMENTAS", "chave": "Competências-chave",
           "ferr": "Ferramentas", "idiomas": "Idiomas", "result": "Resultados"},
    "en": {"resumo": "PROFESSIONAL SUMMARY", "exp": "PROFESSIONAL EXPERIENCE", "form": "EDUCATION",
           "cert": "CERTIFICATIONS", "comp": "SKILLS & TOOLS", "chave": "Core Skills",
           "ferr": "Tools", "idiomas": "Languages", "result": "Results"},
}
NAVY, CINZA = "1F4E79", "555555"


class KitError(RuntimeError):
    pass


# ---------------------------------------------------------------------------
# Clientes: Claude (real) e simulado (sem rede), com a mesma interface
# ---------------------------------------------------------------------------

class KitClient:
    def __init__(self, model: str = "claude-opus-5-5", esforco: str = "medium"):
        import anthropic

        self._anthropic = anthropic
        self.model = model
        self.esforco = esforco
        chave = _ler_env("ANTHROPIC_API_KEY")
        # sem chave, o SDK ainda tenta um login feito com `ant auth login`
        self.client = anthropic.Anthropic(api_key=chave, max_retries=4) if chave else \
            anthropic.Anthropic(max_retries=4)

    def gerar(self, instrucoes: str, mensagem: str, schema: dict, esforco: str | None = None) -> tuple[dict, dict]:
        """Devolve (resposta no formato do schema, uso). As instruções (regras +
        CV + preferências) são iguais em todas as chamadas, então ficam em cache.
        O formato de saída entra no prefixo, então checagem e kit têm caches
        separados: dentro de uma rodada, cada um paga a escrita uma vez e as
        vagas seguintes só a leitura."""
        a = self._anthropic
        try:
            with self.client.beta.messages.stream(
                model=self.model,
                max_tokens=32000,
                # se o filtro de segurança recusar, o próprio servidor refaz em outro modelo
                betas=["server-side-fallback-2026-07-01"],
                fallbacks="default",
                output_config={
                    "effort": esforco or self.esforco,
                    "format": {"type": "json_schema", "schema": schema},
                },
                system=[{"type": "text", "text": instrucoes, "cache_control": {"type": "ephemeral"}}],
                messages=[{"role": "user", "content": mensagem}],
            ) as stream:
                resp = stream.get_final_message()
        except a.AuthenticationError as e:
            raise KitError("401 — confira a ANTHROPIC_API_KEY no .env (ver .env.example).") from e
        except a.APIStatusError as e:
            raise KitError(f"a API do Claude respondeu {e.status_code}: {e.message}") from e
        except a.APIConnectionError as e:
            raise KitError(f"sem conexão com a API do Claude: {e}") from e
        except TypeError as e:
            # o SDK não acha credencial nenhuma (nem chave nem `ant auth login`)
            raise KitError("nenhuma ANTHROPIC_API_KEY encontrada — veja .env.example "
                           "ou rode com --mock.") from e

        if resp.stop_reason == "refusal":
            raise KitError("o modelo recusou gerar este kit (stop_reason=refusal).")
        if resp.stop_reason == "max_tokens":
            raise KitError("a resposta foi cortada no limite de tokens — rode de novo.")
        texto = next((b.text for b in resp.content if b.type == "text"), "")
        try:
            dados = json.loads(texto)
        except json.JSONDecodeError as e:
            raise KitError(f"resposta fora do formato esperado: {texto[:200]}") from e

        u = resp.usage
        uso = {
            "modelo": resp.model,
            "entrada": u.input_tokens,
            "saida": u.output_tokens,
            "cache_leitura": u.cache_read_input_tokens or 0,
            "cache_escrita": u.cache_creation_input_tokens or 0,
        }
        return dados, uso


class MockKitClient:
    """Não chama rede. A checagem não acha lacuna; o kit usa os bullets do
    próprio CV, só pra testar o resto do fluxo (pastas, docx, PDF, checagens)."""

    model = "kit-mock"

    def __init__(self, cv_text: str, bullets_por_experiencia: int = 3):
        self.cv_text = cv_text
        self.n = bullets_por_experiencia

    def gerar(self, instrucoes: str, mensagem: str, schema: dict, esforco: str | None = None) -> tuple[dict, dict]:
        uso = {"modelo": self.model, "entrada": 0, "saida": 0, "cache_leitura": 0, "cache_escrita": 0}
        if "obrigatorios_faltando" in schema["properties"]:
            return {"obrigatorios_faltando": []}, uso
        exps = [{**e, "local": "São Paulo, Brasil", "linha_extra": "", "resultados": "Resultado simulado.",
                 "bullets": (e["bullets"] * self.n)[:self.n]} for e in experiencias_do_cv(self.cv_text)]
        titulo = re.search(r"^Vaga: (.+)$", mensagem, re.M)
        kit = {
            "idioma": "pt",
            "variante": "padrao",
            "titulo_cv": cabecalho_cv(self.cv_text)["titulo"],
            "resumo": "Resumo simulado.",
            "experiencias": exps,
            "formacao": [{"titulo": "Administração de Empresas", "linhas": []}],
            "certificacoes": [],
            "competencias_chave": ["Revenue Operations"],
            "ferramentas": ["Excel (Avançado) · SQL (Básico)"],
            "idiomas": ["Português: Nativo"],
            "carta": f"Carta simulada para {titulo.group(1) if titulo else 'a vaga'}.",
            "respostas": [{"pergunta": "Qual a sua pretensão salarial?",
                           "resposta": "[sua pretensão salarial]"}],
            "palavras_chave": ["revenue operations"],
        }
        return kit, uso


# ---------------------------------------------------------------------------
# Leitura do CV, das regras e das vagas
# ---------------------------------------------------------------------------

def cabecalho_cv(cv_text: str) -> dict:
    """As 3 primeiras linhas do cv.txt: nome, título e contatos. Vão pro .docx
    exatamente como estão — nunca passam pelo modelo."""
    linhas = [l.strip() for l in cv_text.splitlines() if l.strip()]
    return {"nome": linhas[0], "titulo": linhas[1] if len(linhas) > 1 else "",
            "contato": linhas[2] if len(linhas) > 2 else ""}


def experiencias_do_cv(cv_text: str) -> list[dict]:
    """Experiências no formato "Empresa  |  Cargo" + linha do período + bullets.
    Usado pelo cliente simulado."""
    def _cabecalho(i: int) -> bool:
        return "|" in linhas[i] and i + 1 < len(linhas) and bool(re.match(r"\d{2}/\d{4}", linhas[i + 1]))

    exps: list[dict] = []
    linhas = [l.strip() for l in cv_text.splitlines()]
    i = 0
    while i < len(linhas):
        if _cabecalho(i):
            empresa, cargo = (p.strip() for p in linhas[i].split("|", 1))
            exp = {"empresa": empresa, "cargo": cargo,
                   "periodo": linhas[i + 1].split("·")[0].split("(")[0].strip(), "bullets": []}
            i += 2
            while i < len(linhas) and linhas[i] and not linhas[i].isupper() and not _cabecalho(i):
                exp["bullets"].append(linhas[i].lstrip("- ").strip())
                i += 1
            exps.append(exp)
            continue
        i += 1
    return exps


def carregar_regras() -> str:
    return REGRAS_PATH.read_text(encoding="utf-8").strip() if REGRAS_PATH.exists() else ""


def termos_proibidos(regras: str) -> list[str]:
    """Termos das seções "O que ele NÃO tem" e "Inferências que já removemos" do
    regras-cv.md: um item por bullet, separado em "," "/" e " e "; quando o item
    tem um trecho entre aspas, só o trecho conta ("recompra", "1P")."""
    termos: list[str] = []
    for secao in ("O que ele NÃO tem", "Inferências que já removemos"):
        m = re.search(rf"^###[^\n]*{re.escape(secao)}[^\n]*\n(.*?)(?=^#|\Z)", regras, re.M | re.S)
        if not m:
            continue
        for linha in m.group(1).splitlines():
            if not linha.startswith("- "):
                continue
            item = linha[2:].strip()
            aspas = re.findall(r'"([^"]+)"', item)
            partes = aspas or re.split(r",|/| e |\(|\)", item)
            termos += [p.strip() for p in partes if len(p.strip()) >= 2]
    return list(dict.fromkeys(termos))


def ultimo_relatorio(out_dir: Path) -> Path:
    relatorios = sorted(out_dir.glob("vagas_triadas_*.csv"))
    if not relatorios:
        sys.exit(f"Nenhum relatório em {out_dir}/ — rode `python screen.py` primeiro.")
    return relatorios[-1]


def carregar_vagas(out_dir: Path) -> list[dict]:
    """Linhas do último relatório, com a descrição completa da coleta."""
    with ultimo_relatorio(out_dir).open(encoding="utf-8-sig") as f:
        vagas = list(csv.DictReader(f))
    coletadas_path = out_dir / "vagas_coletadas.json"
    descricoes = {}
    if coletadas_path.exists():
        descricoes = {j.get("link"): j.get("descricao", "")
                      for j in json.loads(coletadas_path.read_text(encoding="utf-8"))}
    for n, v in enumerate(vagas, 1):
        v["n"] = n
        v["descricao"] = descricoes.get(v.get("link"), "")
    return vagas


def _sem_acento(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()


def _slug(texto: str, tamanho: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", _sem_acento(texto)).strip("-")[:tamanho].strip("-") or "x"


def pasta_kit(kits_dir: Path, vaga: dict) -> Path:
    # o hash do link separa duas vagas de mesmo título na mesma empresa
    h = hashlib.md5(vaga.get("link", "").encode()).hexdigest()[:6]
    return kits_dir / f"{_slug(vaga['empresa'], 25)}_{_slug(vaga['titulo'])}_{h}"


_PALAVRAS_FORA = {"de", "da", "do", "das", "dos", "e", "em", "para", "com", "a", "o",
                  "the", "of", "and", "for", "in", "to", "at"}
_SUFIXOS_EMPRESA = {"ltda", "sa", "s/a", "eireli", "me", "inc", "llc", "ltd"}


def _camel(texto: str, fora: set[str]) -> str:
    texto = re.sub(r"\((a|o|as|os)\)", "", texto, flags=re.I)  # "Coordenador(a)"
    palavras = re.findall(r"[A-Za-z0-9]+", _sem_acento(texto))
    return "".join(p if p[1:] != p[1:].lower() else p.capitalize()  # mantém "RevOps", "BTG"
                   for p in palavras if p.lower() not in fora)


def nome_pdf(primeiro_nome: str, vaga: dict) -> str:
    """CV_Michel_<Vaga><Empresa>.pdf, sem espaços nem "_" entre as palavras —
    padrão do regras-cv.md. Do título sai o que vem depois de " - ", "|" ou
    "(" (local, modelo de trabalho, tradução)."""
    titulo = re.sub(r"\((a|o|as|os)\)", "", vaga["titulo"], flags=re.I)
    # "(Senior)", "(Pleno)" é parte do cargo — sai dos parênteses antes do corte
    titulo = re.sub(r"\((s[eê]nior|sr|pleno|pl|j[uú]nior|jr|mid-level|lead|i{1,3})\)", r" \1", titulo,
                    flags=re.I)
    titulo = re.split(r"\s[-–—]\s|\||\(|,", titulo)[0]
    empresa = vaga["empresa"].replace(".", "")
    return (f"CV_{_camel(primeiro_nome.capitalize(), set())}_{_camel(titulo, _PALAVRAS_FORA)}"
            f"{_camel(empresa, _SUFIXOS_EMPRESA)}.pdf")


def mensagem_da_vaga(vaga: dict, faltando: list[dict] | None = None) -> str:
    campos = [
        ("Vaga", vaga.get("titulo")),
        ("Empresa", vaga.get("empresa")),
        ("Local", ", ".join(x for x in (vaga.get("cidade"), vaga.get("estado")) if x)),
        ("Remoto", vaga.get("remoto")),
        ("Salário", vaga.get("salario")),
    ]
    texto = "\n".join(f"{k}: {v}" for k, v in campos if v)
    texto += f"\n\nDescrição da vaga:\n{vaga['descricao']}\n\n"
    if faltando is None:
        return texto + PEDIDO_CHECAGEM
    if faltando:
        texto += ("Requisitos obrigatórios que ele não atende (ele decidiu montar mesmo assim; "
                  "não os afirme no CV):\n" + "\n".join(f"- {f['requisito']}" for f in faltando) + "\n\n")
    return texto + "Monte o kit de candidatura para esta vaga."


def montar_instrucoes(cv_text: str, regras: str, preferencias: str) -> str:
    texto = f"{INSTRUCOES}\n<curriculo_mestre>\n{cv_text.strip()}\n</curriculo_mestre>\n"
    if regras:
        texto += f"\n<regras>\n{regras}\n</regras>\n"
    if preferencias:
        texto += f"\n<preferencias>\n{preferencias.strip()}\n</preferencias>\n"
    return texto


# ---------------------------------------------------------------------------
# Checagens e ajustes determinísticos
# ---------------------------------------------------------------------------

def _numeros(texto: str) -> set[str]:
    # "2.000+", "R$300 mil", "~30%", "2024" -> "2000", "300", "30", "2024"
    return {re.sub(r"[.,]", "", m) for m in re.findall(r"\d+(?:[.,]\d+)*", texto)}


def _textos_do_cv(kit: dict) -> list[str]:
    textos = [kit["titulo_cv"], kit["resumo"], *kit["competencias_chave"], *kit["ferramentas"], *kit["idiomas"]]
    for e in kit["experiencias"]:
        textos += [e["empresa"], e["cargo"], e["periodo"], e["linha_extra"], e["resultados"], *e["bullets"]]
    for f in kit["formacao"] + kit["certificacoes"]:
        textos += [f["titulo"], *f["linhas"]]
    return textos


def numeros_nao_encontrados(kit: dict, fontes: str) -> list[str]:
    """Trechos do kit com número que não aparece no cv.txt nem no regras-cv.md.
    Não prova que houve invenção (o modelo pode ter escrito "meio milhão" ou
    somado dois números), mas é onde vale olhar primeiro."""
    conhecidos = _numeros(fontes)
    suspeitos = []
    for texto in _textos_do_cv(kit) + [kit["carta"], *(r["resposta"] for r in kit["respostas"])]:
        faltando = sorted(_numeros(texto) - conhecidos)
        if faltando:
            suspeitos.append(f"{', '.join(faltando)} — \"{texto[:140]}\"")
    return suspeitos


def termos_no_cv(kit: dict, termos: list[str]) -> list[str]:
    """Termos proibidos pelas regras que apareceram no CV ou na carta (as
    respostas ficam de fora: lá "não tenho Python" é uma resposta honesta)."""
    texto = " ".join(_textos_do_cv(kit) + [kit["carta"]])
    return [t for t in termos if re.search(rf"(?<!\w){re.escape(t)}(?!\w)", texto, re.I)]


def marcadores(kit: dict) -> list[str]:
    """[sua pretensão salarial] e afins, que você precisa preencher."""
    achados = []
    for texto in [kit["carta"], *(r["resposta"] for r in kit["respostas"])]:
        achados += re.findall(r"\[[^\]]{3,80}\]", texto)
    return list(dict.fromkeys(achados))


def titulo_em_uma_linha(titulo: str) -> str:
    """Corta frentes do fim (no " | ") até caber em uma linha."""
    partes = [p.strip() for p in titulo.split("|")]
    while len(partes) > 1 and len(" | ".join(partes)) > TITULO_MAX_CHARS:
        partes.pop()
    return " | ".join(partes)


def cortar_um_bullet(kit: dict) -> bool:
    """Tira o último bullet (o menos relevante) da experiência com mais bullets;
    no empate, da mais antiga. Nenhuma fica com menos de 2, e a linha de
    resultados nunca sai. Depois disso, corta competências-chave (mínimo 4)."""
    candidatas = [e for e in kit["experiencias"] if len(e["bullets"]) > 2]
    if candidatas:
        maior = max(len(e["bullets"]) for e in candidatas)
        [e for e in candidatas if len(e["bullets"]) == maior][-1]["bullets"].pop()
        return True
    if len(kit["competencias_chave"]) > 4:
        kit["competencias_chave"].pop()
        return True
    return False


# ---------------------------------------------------------------------------
# Arquivos de saída
# ---------------------------------------------------------------------------

def escrever_docx(kit: dict, cab: dict, path: Path) -> None:
    """Mesmo template dos CVs anteriores (regras-cv.md, "Template visual")."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor, Twips

    rot = ROTULOS.get(kit["idioma"], ROTULOS["pt"])
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Twips(900), Twips(900)
    s.left_margin, s.right_margin = Twips(1100), Twips(1100)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    def par(antes=0.0, depois=0.0, centro=False, justificado=False, estilo=None):
        p = doc.add_paragraph(style=estilo)
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(depois)
        if centro:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif justificado:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def run(p, texto, negrito=False, italico=False, tamanho=None, cor=None):
        r = p.add_run(texto)
        r.bold, r.italic = negrito, italico
        if tamanho:
            r.font.size = Pt(tamanho)
        if cor:
            r.font.color.rgb = RGBColor.from_string(cor)
        return r

    def secao(titulo):
        p = par(antes=9, depois=4)
        run(p, titulo, negrito=True, tamanho=11.5, cor=NAVY)
        borda = OxmlElement("w:pBdr")
        baixo = OxmlElement("w:bottom")
        for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "A6A6A6"}.items():
            baixo.set(qn(k), v)
        borda.append(baixo)
        p._p.get_or_add_pPr().append(borda)

    def bullet(texto):
        p = par(depois=1, estilo="List Bullet")
        # recuo dos CVs anteriores: bolinha a ~0,45 cm da margem, texto a ~0,85 cm.
        # O estilo de lista traz uma tabulação em 360 twips que puxaria a 1ª linha
        # pra perto da bolinha; troca por uma no início do texto.
        pf = p.paragraph_format
        pf.left_indent = Cm(0.85)
        pf.first_line_indent = Cm(-0.4)
        pf.tab_stops.add_tab_stop(Twips(360), WD_TAB_ALIGNMENT.CLEAR)
        pf.tab_stops.add_tab_stop(Cm(0.85))
        run(p, texto)

    run(par(centro=True), cab["nome"], negrito=True, tamanho=17)
    run(par(antes=2, centro=True), kit["titulo_cv"], negrito=True, tamanho=11.5, cor=NAVY)
    run(par(antes=2, centro=True), cab["contato"], cor=CINZA)

    secao(rot["resumo"])
    run(par(justificado=True), kit["resumo"])

    secao(rot["exp"])
    for i, e in enumerate(kit["experiencias"]):
        run(par(antes=0 if i == 0 else 6), f"{e['empresa']} | {e['cargo']}", negrito=True, tamanho=11.5)
        run(par(), " · ".join(x for x in (e["periodo"], e["local"]) if x), italico=True, cor=CINZA)
        if e["linha_extra"]:
            run(par(), e["linha_extra"], cor=CINZA)
        for b in e["bullets"]:
            bullet(b)
        if e["resultados"]:
            bullet(f"{rot['result']}: {e['resultados']}")

    def itens_com_linhas(itens):
        for i, f in enumerate(itens):
            run(par(antes=0 if i == 0 else 4), f["titulo"], negrito=True)
            for linha in f["linhas"]:
                run(par(), linha, italico=True, cor=CINZA)

    secao(rot["form"])
    itens_com_linhas(kit["formacao"])
    if kit["certificacoes"]:
        secao(rot["cert"])
        itens_com_linhas(kit["certificacoes"])

    secao(rot["comp"])
    for i, (rotulo, itens) in enumerate(((rot["chave"], kit["competencias_chave"]),
                                         (rot["ferr"], kit["ferramentas"]),
                                         (rot["idiomas"], kit["idiomas"]))):
        run(par(antes=0 if i == 0 else 4, depois=1), rotulo, negrito=True)
        for item in itens:
            bullet(item)
    doc.save(path)


class ConversorWord:
    """Word (via COM) contando páginas e salvando o PDF — o mesmo motor que vai
    abrir o arquivo do outro lado. Uma instância separada (DispatchEx), pra não
    mexer em documentos que você tenha aberto."""

    def __enter__(self):
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()  # o painel chama de uma thread própria
        self.word = win32com.client.DispatchEx("Word.Application")
        self.word.Visible = False
        self.word.DisplayAlerts = 0
        return self

    def __exit__(self, *exc):
        import pythoncom

        self.word.Quit()
        pythoncom.CoUninitialize()

    def _abrir(self, docx: Path):
        return self.word.Documents.Open(str(docx.resolve()), ReadOnly=True, AddToRecentFiles=False)

    def paginas(self, docx: Path) -> int:
        doc = self._abrir(docx)
        try:
            doc.Repaginate()
            return doc.ComputeStatistics(2)  # wdStatisticPages
        finally:
            doc.Close(SaveChanges=0)

    def pdf(self, docx: Path, destino: Path) -> None:
        doc = self._abrir(docx)
        try:
            doc.SaveAs2(str(destino.resolve()), FileFormat=17)  # wdFormatPDF
        finally:
            doc.Close(SaveChanges=0)


def encaixar_em_2_paginas(kit: dict, cab: dict, docx: Path, conversor) -> int:
    """Escreve o docx e corta bullets até ter no máximo 2 páginas. Devolve o
    número de páginas final (1 quer dizer que faltou conteúdo)."""
    while True:
        escrever_docx(kit, cab, docx)
        paginas = conversor.paginas(docx)
        if paginas <= PAGINAS or not cortar_um_bullet(kit):
            return paginas


def escrever_md(kit: dict, vaga: dict, arquivo: str, paginas: int | None, faltando: list[dict],
                suspeitos: list[str], proibidos: list[str], path: Path) -> None:
    linhas = [
        f"# {vaga['titulo']} — {vaga['empresa']}",
        "",
        f"- Link: {vaga.get('link', '')}",
        f"- Nota do Jev: {vaga.get('nota_fit', '?')} ({vaga.get('recomendacao', '?')})",
        f"- Variante: {kit['variante']} · idioma: {kit['idioma']}",
        f"- Gerado em {dt.date.today().isoformat()}",
        "",
        "## Antes de enviar",
        "",
        f"- [ ] Abrir e revisar o `{arquivo}`",
    ]
    if paginas is not None and paginas != PAGINAS:
        linhas.append(f"- [ ] O CV ficou com {paginas} página(s), não {PAGINAS} — ajustar à mão")
    linhas += [f"- [ ] Preencher {m}" for m in marcadores(kit)]
    if proibidos:
        linhas.append(f"- [ ] Tirar o que as regras proíbem: {', '.join(proibidos)}")
    if suspeitos:
        linhas.append("- [ ] Conferir números que não aparecem no cv.txt nem no regras-cv.md:")
        linhas += [f"    - {s}" for s in suspeitos]
    linhas += ["", "## Requisitos obrigatórios que você não atende", ""]
    linhas += [f"- {f['requisito']} — {f['motivo']}" for f in faltando] or ["- nenhum"]
    linhas += ["", "## Carta", "", kit["carta"], "", "## Respostas para o formulário", ""]
    for r in kit["respostas"]:
        linhas += [f"**{r['pergunta']}**", "", r["resposta"], ""]
    linhas += ["## Palavras-chave da vaga que o CV cobre", "", ", ".join(kit["palavras_chave"])]
    path.write_text("\n".join(linhas) + "\n", encoding="utf-8")


REGISTRO_CAMPOS = ["data", "empresa", "vaga", "fonte", "link", "nota_jev", "lacunas", "decisao",
                   "variante", "arquivo", "paginas"]


def registrar(kits_dir: Path, vaga: dict, faltando: list[dict], decisao: str, variante: str = "",
              arquivo: str = "", paginas: int | str = "") -> None:
    """Uma linha por vaga processada em output/kits/registro.csv (o "log" das regras)."""
    path = kits_dir / "registro.csv"
    novo = not path.exists()
    kits_dir.mkdir(parents=True, exist_ok=True)
    with path.open("a", newline="", encoding="utf-8-sig" if novo else "utf-8") as f:
        w = csv.DictWriter(f, fieldnames=REGISTRO_CAMPOS)
        if novo:
            w.writeheader()
        w.writerow({"data": dt.date.today().isoformat(), "empresa": vaga["empresa"], "vaga": vaga["titulo"],
                    "fonte": vaga.get("fonte", ""), "link": vaga.get("link", ""),
                    "nota_jev": vaga.get("nota_fit", ""), "lacunas": "; ".join(f["requisito"] for f in faltando),
                    "decisao": decisao, "variante": variante, "arquivo": arquivo, "paginas": paginas})


def _custo(usos: list[dict]) -> float | None:
    total = 0.0
    for uso in usos:
        p = PRECO_MTOK.get(uso["modelo"])
        if not p:
            return None
        total += sum(uso[k] * p[k] for k in p) / 1_000_000
    return total


def perguntar_no_terminal(pergunta: str) -> bool | None:
    """True/False = respondeu; None = não tem quem responda (rodando agendado,
    sem terminal) — aí a vaga fica para decidir no painel."""
    if not sys.stdin.isatty():
        return None
    try:
        return input(f"{pergunta} [s/N] ").strip().lower() in ("s", "sim", "y")
    except EOFError:  # terminal que se diz interativo mas não tem teclado ligado
        print()
        return None


# ---------------------------------------------------------------------------
# Fluxo
# ---------------------------------------------------------------------------

def completar_descricao(vaga: dict) -> dict:
    """Vagas do beBee chegam com ~300 caracteres; busca o texto inteiro."""
    if vaga.get("fonte") != "bebee" or len(vaga.get("descricao") or "") >= DESCRICAO_CURTA:
        return vaga
    from collectors.bebee import fetch_bebee_descricao

    try:
        completa = fetch_bebee_descricao(vaga["link"])
    except requests.RequestException:
        completa = ""
    if len(completa) > len(vaga.get("descricao") or ""):
        return {**vaga, "descricao": completa}
    print("  [aviso] não consegui a descrição completa no beBee — o kit sai com o resumo.", file=sys.stderr)
    return vaga


def gerar_kit(client, cv_text: str, regras: str, preferencias: str, vaga: dict, kits_dir: Path,
              conversor=None, pasta_pdf: Path | None = None, refazer: bool = False,
              perguntar=perguntar_no_terminal, faltando: list[dict] | None = None) -> Path | None:
    """Gera o kit de uma vaga. Devolve a pasta do kit, ou None se a vaga foi
    pulada ou ficou para decidir. `conversor` é o ConversorWord (None = só
    .docx); `perguntar(texto)` decide "montar mesmo assim?" e "substituir o
    PDF?" (None = sem resposta). `faltando`: lacunas já checadas e aceitas
    (decisão "montar mesmo assim" no painel) — pula a checagem."""
    estado = Candidaturas(kits_dir / "candidaturas.json")
    pasta = pasta_kit(kits_dir, vaga)
    if (pasta / "kit.json").exists() and not refazer:
        print(f"  já existe: {pasta} (use --refazer pra montar um novo)")
        return pasta

    vaga = completar_descricao(vaga)
    if not vaga.get("descricao"):
        raise KitError("vaga sem descrição em output/vagas_coletadas.json — rode `python screen.py` "
                       "de novo para coletar as vagas de hoje.")

    instrucoes = montar_instrucoes(cv_text, regras, preferencias)
    usos = []

    # 1) requisitos obrigatórios: se faltar algum, pergunta antes de gastar com o kit
    if faltando is None:
        checagem, uso = client.gerar(instrucoes, mensagem_da_vaga(vaga), CHECAGEM_SCHEMA, esforco="low")
        usos.append(uso)
        faltando = checagem["obrigatorios_faltando"]
        if faltando:
            print("  Requisitos obrigatórios que você não atende:")
            for f in faltando:
                print(f"    - {f['requisito']} — {f['motivo']}")
            resposta = perguntar("  Montar mesmo assim?")
            if resposta is None:
                print("  ficou em \"Para decidir\" no painel (python painel.py).")
                estado.atualizar(vaga, "decidir", lacunas=faltando, descricao=vaga["descricao"])
                return None
            if not resposta:
                print("  pulada.")
                registrar(kits_dir, vaga, faltando, "pulada")
                estado.atualizar(vaga, "pulada", lacunas=faltando, descricao=vaga["descricao"])
                return None

    # 2) o kit
    kit, uso = client.gerar(instrucoes, mensagem_da_vaga(vaga, faltando), KIT_SCHEMA)
    usos.append(uso)
    kit["titulo_cv"] = titulo_em_uma_linha(kit["titulo_cv"])
    if kit["variante"] != "financeira":
        kit["certificacoes"] = []  # regras: seção de certificações só na variante financeira

    cab = cabecalho_cv(cv_text)
    arquivo = nome_pdf(cab["nome"].split()[0], vaga)
    pasta.mkdir(parents=True, exist_ok=True)
    docx = pasta / arquivo.replace(".pdf", ".docx")
    pdf = pasta / arquivo
    paginas = None
    if conversor:
        paginas = encaixar_em_2_paginas(kit, cab, docx, conversor)
        conversor.pdf(docx, pdf)
    else:
        escrever_docx(kit, cab, docx)

    suspeitos = numeros_nao_encontrados(kit, f"{cv_text}\n{regras}")
    proibidos = termos_no_cv(kit, termos_proibidos(regras))
    escrever_md(kit, vaga, pdf.name if conversor else docx.name, paginas, faltando, suspeitos, proibidos,
                pasta / "kit.md")

    copia = None
    if conversor and pasta_pdf:
        if not pasta_pdf.is_dir():
            print(f"  [aviso] pasta {pasta_pdf} não existe — o PDF ficou só em {pasta}", file=sys.stderr)
        elif (pasta_pdf / arquivo).exists() and not refazer and \
                not perguntar(f"  Já existe {arquivo} em {pasta_pdf}. Substituir?"):
            print(f"  mantive o PDF antigo da pasta; o novo está em {pdf}")
        else:
            copia = shutil.copy2(pdf, pasta_pdf / arquivo)

    dados = {
        "vaga": {k: vaga.get(k, "") for k in ("titulo", "empresa", "link", "fonte", "nota_fit", "recomendacao")},
        "gerado_em": dt.datetime.now().isoformat(timespec="seconds"),
        "uso": usos, "paginas": paginas, "pdf": str(pdf) if conversor else "",
        "pdf_copiado": str(copia or ""),
        "obrigatorios_faltando": faltando, "numeros_a_conferir": suspeitos, "termos_proibidos": proibidos,
        "kit": kit,
    }
    (pasta / "kit.json").write_text(json.dumps(dados, ensure_ascii=False, indent=2), encoding="utf-8")
    registrar(kits_dir, vaga, faltando, "montada", kit["variante"], arquivo, paginas or "")
    custo = _custo(usos)
    atual = next((e for e in estado.todas() if e["link"] == vaga["link"]), {})
    estado.atualizar(vaga, None if atual.get("status") in POS_ENVIO else "pronta",
                     pasta=str(pasta), pdf=dados["pdf"], lacunas=faltando, alertas=alertas_do_kit(dados),
                     gerado_em=dados["gerado_em"], custo=custo, descricao=vaga["descricao"], erro="")

    print(f"  ok: {copia or pasta}" + (f"  (~US$ {custo:.3f})" if custo is not None else ""))
    if paginas is not None and paginas != PAGINAS:
        print(f"  [atenção] o CV ficou com {paginas} página(s) — veja o kit.md")
    if proibidos:
        print(f"  [atenção] termos que as regras proíbem: {', '.join(proibidos)} — veja o kit.md")
    if suspeitos:
        print(f"  [atenção] {len(suspeitos)} trecho(s) com número que não está no cv.txt/regras — veja o kit.md")
    return pasta


def gerar_sem_terminal(config: dict, vaga: dict, kits_dir: Path, mock: bool = False, refazer: bool = False,
                       faltando: list[dict] | None = None) -> Path | None:
    """gerar_kit() com tudo lido do config, para quem não tem terminal (o painel
    e o sincronizar.py): lacuna obrigatória vai para "Para decidir" e um PDF de
    mesmo nome já na pasta só é substituído quando é para refazer."""
    cv_text = Path(config["cv_path"]).read_text(encoding="utf-8")
    cfg_kit = config.get("kit") or {}
    client = MockKitClient(cv_text) if mock else KitClient(
        model=cfg_kit.get("modelo", "claude-opus-5-5"), esforco=cfg_kit.get("esforco", "medium"))
    pasta_pdf = None if mock or not cfg_kit.get("pasta_pdf") else Path(cfg_kit["pasta_pdf"]).expanduser()
    try:
        conversor_ctx = ConversorWord()
        conversor = conversor_ctx.__enter__()
    except Exception as e:  # noqa: BLE001 — sem Word (ou sem pywin32) ainda dá pra gerar o .docx
        print(f"  [aviso] sem Word ({e}) — o CV sai só em .docx.", file=sys.stderr)
        conversor_ctx = conversor = None
    try:
        return gerar_kit(client, cv_text, carregar_regras(), carregar_preferencias(config), vaga, kits_dir,
                         conversor=conversor, pasta_pdf=pasta_pdf, refazer=refazer,
                         perguntar=lambda _pergunta: None, faltando=faltando)
    finally:
        if conversor_ctx:
            conversor_ctx.__exit__(None, None, None)


def listar(vagas: list[dict], kits_dir: Path, todas: bool, busca: str | None) -> None:
    sel = vagas if todas else [v for v in vagas if v["recomendacao"] == "aplicar_agora"]
    if busca:
        b = busca.lower()
        sel = [v for v in sel if b in f"{v['empresa']} {v['titulo']}".lower()]
    if not sel:
        print("Nenhuma vaga com esse filtro.")
        return
    status = {e["link"]: e.get("status", "") for e in Candidaturas(kits_dir / "candidaturas.json").todas()}
    marca = {"decidir": "?", "pulada": "-"}
    mostrar = sel if busca else sel[:MAX_NA_LISTA]
    for v in mostrar:
        s = status.get(v["link"], "")
        feito = marca.get(s, "✓" if s else " ")
        print(f"  {v['n']:4d} {feito} {float(v['nota_fit'] or 0):4.1f}  {v['titulo']} — {v['empresa']}")
    if len(sel) > len(mostrar):
        print(f"  ... +{len(sel) - len(mostrar)} (use --busca para filtrar)")
    print("\nGere com: python kit.py <número> [<número> ...]   "
          "(✓ = kit já gerado, ? = para decidir no painel, - = pulada)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("numeros", nargs="*", type=int, help="números das vagas na lista")
    ap.add_argument("--link", help="gera o kit da vaga com este link")
    ap.add_argument("--busca", help="filtra a lista por empresa/título")
    ap.add_argument("--todas", action="store_true", help="lista todas as recomendações")
    ap.add_argument("--refazer", action="store_true", help="monta de novo mesmo se o kit já existir")
    ap.add_argument("--sempre", action="store_true",
                    help="não pergunta nada: monta mesmo se faltar requisito obrigatório e "
                    "substitui o PDF se já existir um com o mesmo nome na pasta")
    ap.add_argument("--mock", action="store_true", help="kit simulado, sem chamar a API")
    ap.add_argument("--config", default="config.yaml")
    args = ap.parse_args()

    config = yaml.safe_load(Path(args.config).read_text(encoding="utf-8"))
    out_dir = Path(config.get("output_dir", "output"))
    # kits simulados ficam à parte, pra nunca passarem por kit pronto na lista
    kits_dir = out_dir / ("kits_mock" if args.mock else "kits")
    vagas = carregar_vagas(out_dir)

    if args.link:
        escolhidas = [v for v in vagas if v.get("link") == args.link]
        if not escolhidas:
            sys.exit("Esse link não está no último relatório.")
    elif args.numeros:
        por_numero = {v["n"]: v for v in vagas}
        faltando = [n for n in args.numeros if n not in por_numero]
        if faltando:
            sys.exit(f"Número(s) fora da lista: {faltando}")
        escolhidas = [por_numero[n] for n in args.numeros]
    else:
        listar(vagas, kits_dir, args.todas, args.busca)
        return

    cv_text = Path(config["cv_path"]).read_text(encoding="utf-8")
    regras = carregar_regras()
    if not regras:
        print(f"  [aviso] {REGRAS_PATH} não encontrado — o kit sai só com as instruções genéricas.",
              file=sys.stderr)
    preferencias = carregar_preferencias(config)
    cfg_kit = config.get("kit") or {}
    if args.mock:
        client = MockKitClient(cv_text)
    else:
        client = KitClient(model=cfg_kit.get("modelo", "claude-opus-5-5"),
                           esforco=cfg_kit.get("esforco", "medium"))
    # o PDF de teste (--mock) não vai pra pasta de CVs de verdade
    pasta_pdf = None if args.mock else Path(cfg_kit.get("pasta_pdf", "")).expanduser() \
        if cfg_kit.get("pasta_pdf") else None
    perguntar = (lambda _p: True) if args.sempre else perguntar_no_terminal

    try:
        conversor_ctx = ConversorWord()
        conversor = conversor_ctx.__enter__()
    except Exception as e:  # noqa: BLE001 — sem Word (ou sem pywin32) ainda dá pra gerar o .docx
        print(f"  [aviso] não consegui abrir o Word ({e}) — o CV sai só em .docx, sem PDF.", file=sys.stderr)
        conversor_ctx = conversor = None

    falhas = 0
    try:
        for v in escolhidas:
            print(f"[{v['n']}] {v['titulo']} — {v['empresa']}")
            try:
                gerar_kit(client, cv_text, regras, preferencias, v, kits_dir, conversor=conversor,
                          pasta_pdf=pasta_pdf, refazer=args.refazer, perguntar=perguntar)
            except KitError as e:
                falhas += 1
                print(f"  [erro] {e}", file=sys.stderr)
    finally:
        if conversor_ctx:
            conversor_ctx.__exit__(None, None, None)
    if falhas:
        sys.exit(1)


if __name__ == "__main__":
    main()
